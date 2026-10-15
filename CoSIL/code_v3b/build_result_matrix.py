"""Maintain a 300-issue matrix and append experiment columns by instance_id."""
import argparse
import csv
from datetime import datetime
import hashlib
import html
import json
from pathlib import Path
import re
import sys
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
STORE = ROOT / 'results/comparison/逐题定位对照数据.json'
OUTPUT = ROOT / 'docs/qwen-300'
TITLE = '跨实验逐题定位对照表'


def sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def normalize(name):
    return name.strip().strip('`').split('(', 1)[0].split('#', 1)[0].strip().rsplit('.', 1)[-1]


def predictions(row):
    """Preserve displayed names, with the existing scorer's file+leaf deduplication."""
    found = []
    seen = set()
    for file, values in row.get('found_related_locs', {}).items():
        if not isinstance(values, (str, list)):
            raise ValueError(f'Unsupported locations for {row["instance_id"]}: {file}')
        for text in [values] if isinstance(values, str) else values:
            if not isinstance(text, str):
                raise ValueError('Location entries must be strings')
            for line in text.replace('\\n', '\n').splitlines():
                if not line.strip():
                    continue
                match = re.match(r'^\s*(function|class|variable)\s*:\s*(.+)', line)
                if not match:
                    raise ValueError(f'Unsupported location label: {line!r}')
                kind, name = match[1], match[2].strip()
                key = (file, kind, normalize(name))
                if key not in seen:
                    seen.add(key)
                    found.append(dict(file=file, kind=kind, name=name, normalized_name=key[2]))
    return found


def read_experiment(path, ids):
    records = {}
    for number, line in enumerate(path.read_text().splitlines(), 1):
        if not line.strip():
            continue
        row = json.loads(line)
        iid = row['instance_id']
        if iid in records:
            raise ValueError(f'Duplicate instance_id at line {number}: {iid}')
        if iid not in ids:
            raise ValueError(f'Issue outside the fixed benchmark: {iid}')
        failed = row.get('status') == 'failed' or row.get('audit_status') == 'failed'
        if not failed and not isinstance(row.get('found_related_locs'), dict):
            raise ValueError(f'Missing found_related_locs dictionary: {iid}')
        records[iid] = {key: row[key] for key in ('instance_id', 'arm', 'status',
            'audit_status', 'protocol', 'base_commit', 'found_files', 'found_related_locs',
            'error', 'message') if key in row}
        if not failed:
            predictions(records[iid])
    return records


def score(row, gold):
    targets = {(file, 'function', normalize(name)) for file, names in
               gold['gold_functions_by_file'].items() for name in names}
    if row is None:
        return dict(state='missing', positions=[], top1=None, top3=None, top5=None)
    failed = row.get('status') == 'failed' or row.get('audit_status') == 'failed'
    positions = [] if failed else predictions(row)
    keys = [(p['file'], p['kind'], p['normalized_name']) for p in positions]
    values = {f'top{k}': bool(set(keys[:k]) & targets) if targets else None for k in (1, 3, 5)}
    return dict(state='failed' if failed else 'empty' if not positions else 'complete',
                positions=positions, **values)


def label(value):
    return 'NA（无函数真值）' if value is None else '✓ 命中' if value else '✗ 未命中'


def cell_text(result, row):
    if result['state'] == 'missing':
        return '未提供结果'
    if result['state'] == 'failed':
        return '运行失败：' + str(row.get('error', '未知错误'))
    lines = ['函数 Top-5：' + label(result['top5'])]
    if not result['positions']:
        lines.append('无预测位置')
    for index, p in enumerate(result['positions'], 1):
        lines.append(f"{index}. {p['file']} :: {p['kind']}: {p['name']}")
    return '\n'.join(lines)


def relative(path):
    try:
        return str(path.resolve().relative_to(ROOT))
    except ValueError:
        return str(path.resolve())


def new_store(gold_file, ids):
    store = dict(schema_version=1, benchmark='SWE-bench Lite test / CoSIL v3c gold',
                 issue_order=ids, gold_file=relative(gold_file), gold_sha256=sha(gold_file), projects=[])
    for arm in ('on', 'off'):
        path = ROOT / f'results/qwen_300/fixed/{arm}/loc_{arm}.jsonl'
        records = read_experiment(path, set(ids))
        if set(records) != set(ids) or any(row.get('arm') != arm for row in records.values()):
            raise ValueError(f'Need exactly 300 matching CoSIL {arm} cases')
        store['projects'].append(dict(column=f'CoSIL/{arm}', model='Qwen2.5-Coder-7B-Instruct',
            description='正确函数功能说明' if arm == 'on' else '无函数功能说明',
            source_file=relative(path), source_sha256=sha(path), records=records))
    return store


def write_csv(path, header, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def export(store, gold, folder):
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        sys.path.insert(0, str(ROOT / 'code_v3b/.matrix_deps'))
        from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ids = store['issue_order']
    projects = store['projects']
    header = ['issue'] + [p['column'] for p in projects]
    results = {p['column']: {iid: score(p['records'].get(iid), gold[iid]) for iid in ids}
               for p in projects}
    rows = [[iid] + [cell_text(results[p['column']][iid], p['records'].get(iid))
                    for p in projects] for iid in ids]
    folder.mkdir(parents=True, exist_ok=True)
    write_csv(folder / f'{TITLE}.csv', header, rows)

    workbook_path = folder / f'{TITLE}.xlsx'
    managed_names = ['定位结果', 'Top1命中', 'Top3命中', 'Top5命中', '实验信息']
    if workbook_path.exists():
        book = load_workbook(workbook_path)
        for name in managed_names:
            if name in book.sheetnames:
                del book[name]
    else:
        book = Workbook()
        del book[book.active.title]
    sheet = book.create_sheet('定位结果')
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    for index, iid in enumerate(ids, 2):
        count = max(len(results[p['column']][iid]['positions']) for p in projects)
        sheet.row_dimensions[index].height = max(95, min(350, 25 + count * 24))
        for col, p in enumerate(projects, 2):
            value = results[p['column']][iid]
            color = ('E2F0D9' if value['top5'] is True else 'FCE4D6'
                     if value['top5'] is False else 'E7E6E6')
            sheet.cell(index, col).fill = PatternFill('solid', fgColor=color)
    for k in (1, 3, 5):
        ws = book.create_sheet(f'Top{k}命中')
        ws.append(header)
        values = []
        for iid in ids:
            row = [iid]
            for p in projects:
                item = results[p['column']][iid]
                v = item[f'top{k}']
                row.append('未提供' if item['state'] == 'missing' else 'NA' if v is None else int(v))
            values.append(row)
            ws.append(row)
        if k == 5:
            write_csv(folder / '跨实验逐题Top5命中表.csv', header, values)
    info = book.create_sheet('实验信息')
    info.append(['实验列', '模型', '条件', '提供题数', '函数评分题数', 'Top1', 'Top3', 'Top5', '原始结果', 'SHA256'])
    summary = []
    for p in projects:
        data = results[p['column']]
        valid = [value for value in data.values() if value['state'] != 'missing' and value['top5'] is not None]
        hits = {f'top{k}': sum(v[f'top{k}'] is True for v in valid) for k in (1, 3, 5)}
        item = dict(column=p['column'], rows=len(p['records']), function_cases=len(valid), **hits)
        summary.append(item)
        text = [f'{hits[f"top{k}"]}/{len(valid)}' for k in (1, 3, 5)]
        info.append([p['column'], p['model'], p['description'], len(p['records']), len(valid),
                     *text, p['source_file'], p['source_sha256']])
    # Keep supplemental analysis sheets from previous experiments intact.
    managed_sheets = [book[name] for name in managed_names]
    book._sheets = managed_sheets + [ws for ws in book if ws.title not in managed_names]
    for ws in managed_sheets:
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 26
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24476B')
        ws.column_dimensions['A'].width = 39
        for col in range(2, ws.max_column+1):
            ws.column_dimensions[get_column_letter(col)].width = 76 if ws == sheet else 23
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    book.save(workbook_path)

    md = '# 跨实验逐题定位对照表\n\n'
    md += '基准仓库：`CoSIL`。每行一个 issue，每个实验条件占一列，后续项目按 issue ID 追加列。\n\n'
    md += '单元格包含函数 Top-5 判定和全部解析后的预测位置；保留类/方法前缀。排序与去重沿用 v3c 评分规则，类与变量占推荐位置，第六位以后不计入 Top-5。`✓`=命中，`✗`=未命中，NA=没有函数真值；无预测位置、运行失败、未提供结果分别标记。\n\n'
    md += f'[Excel 工作簿]({TITLE}.xlsx) · [完整定位 CSV]({TITLE}.csv) · [Top-5 命中 CSV](跨实验逐题Top5命中表.csv)\n\n'
    md += '| issue | ' + ' | '.join(p['column'].replace('|', '\\|') for p in projects) + ' |\n'
    md += '|---|' + '---|' * len(projects) + '\n'
    for row in rows:
        cells = [html.escape(value).replace('|', '&#124;').replace('\n', '<br>') for value in row]
        md += '| ' + ' | '.join(cells) + ' |\n'
    md += '\n未经归一化去重的原始定位字典和实验列来源保存在 `results/comparison/逐题定位对照数据.json`。\n'
    (folder / f'{TITLE}.md').write_text(md, encoding='utf-8')
    store['last_export'] = dict(time=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
        issues=len(ids), columns=header[1:], summary=summary, output_folder=relative(folder))
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--store', type=Path, default=STORE)
    parser.add_argument('--output-dir', type=Path, default=OUTPUT)
    parser.add_argument('--add', help='New project/condition column, e.g. Agentless/Qwen')
    parser.add_argument('--input', type=Path, help='JSONL with instance_id and found_related_locs')
    parser.add_argument('--model', default='未注明')
    parser.add_argument('--description', default='')
    parser.add_argument('--allow-partial', action='store_true')
    parser.add_argument('--replace', action='store_true', help='Explicitly replace an existing column')
    args = parser.parse_args()
    if bool(args.add) != bool(args.input):
        parser.error('--add and --input must be provided together')
    gold_file = ROOT / 'ground_truth_300.json'
    gold = json.loads(gold_file.read_text())
    ids = json.loads((ROOT / 'results/qwen_300/fixed_run_manifest.json').read_text())['instances']
    if len(ids) != 300 or len(set(ids)) != 300 or set(ids) != set(gold):
        raise ValueError('Expected the same 300 issues and gold labels')
    store = json.loads(args.store.read_text()) if args.store.exists() else new_store(gold_file, ids)
    if store['issue_order'] != ids or store['gold_sha256'] != sha(gold_file):
        raise ValueError('Benchmark/gold changed; create a separately versioned matrix')
    if args.add:
        records = read_experiment(args.input.resolve(), set(ids))
        if not args.allow_partial and set(records) != set(ids):
            raise ValueError(f'Expected 300 results; received {len(records)}. Use --allow-partial to mark missing rows.')
        for iid, row in records.items():
            if row.get('base_commit') and row['base_commit'] != gold[iid]['base_commit']:
                raise ValueError(f'Base commit mismatch: {iid}')
        old = next((i for i,p in enumerate(store['projects']) if p['column'] == args.add), None)
        if old is not None and not args.replace:
            raise ValueError('Column already exists; use --replace for an intentional update')
        project = dict(column=args.add, model=args.model, description=args.description,
            source_file=relative(args.input), source_sha256=sha(args.input), records=records)
        if old is None:
            store['projects'].append(project)
        else:
            store['projects'][old] = project
    summary = export(store, gold, args.output_dir)
    args.store.parent.mkdir(parents=True, exist_ok=True)
    tmp = args.store.with_suffix('.tmp')
    tmp.write_text(json.dumps(store, ensure_ascii=False, indent=2) + '\n')
    tmp.replace(args.store)
    print(json.dumps(dict(issues=300, columns=[p['column'] for p in store['projects']], summary=summary,
                         output_folder=str(args.output_dir)), ensure_ascii=False))


if __name__ == '__main__':
    main()
