"""Frozen-prediction reports and supplemental workbook sheets."""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import statistics
import sys
from datetime import datetime

from qwen_adapt.runner import ROOT,WORK,atomic,sha
from qwen_adapt.evaluate import COSIL,read_jsonl,common_metrics,native_metrics,evaluate_folder
from qwen_adapt.diagnostics import analyze


def rate(value):
    return 'NA' if value is None else f'{value:.2%}'


def metric_line(label, value):
    return '| '+label+' | '+str(value['denominator'])+' | '+' | '.join(rate(value[k]) for k in ('precision','recall','f1','sl'))+' |\n'


def project_hash(project):
    return hashlib.sha256(json.dumps(project,ensure_ascii=False,sort_keys=True).encode()).hexdigest()


def baseline_projects():
    store=json.loads((COSIL/'results/comparison/逐题定位对照数据.json').read_text())
    return {p['column']:p for p in store['projects'] if p['column'] in ('CoSIL/on','CoSIL/off','GraphLocator/Qwen')}


def paired_interval(old,new,ids,field='strict'):
    import numpy as np
    differences=[new[iid][field]['f1']-old[iid][field]['f1'] for iid in ids if old[iid][field] is not None and new[iid][field] is not None]
    if not differences:
        return None
    a=np.array(differences);rng=np.random.default_rng(0)
    samples=a[rng.integers(0,len(a),size=(5000,len(a)))].mean(axis=1)
    return dict(mean=float(a.mean()),low=float(np.quantile(samples,.025)),high=float(np.quantile(samples,.975)),
        positive=sum(x>0 for x in differences),negative=sum(x<0 for x in differences),unchanged=sum(x==0 for x in differences))


def generate():
    inputs=json.loads((WORK/'inputs_manifest.json').read_text())
    ids=inputs['instances']
    manifest=json.loads((WORK/'full/manifest.json').read_text())
    selected=manifest['variants'][0];folder=WORK/'full'/selected
    metrics,details,audits=evaluate_folder(folder,ids)
    if metrics['missing']:
        raise RuntimeError('Full report requires a result or recorded failure for all 300 issue IDs')
    records=read_jsonl(folder/'loc_results.jsonl')
    gold=json.loads((COSIL/'ground_truth_300.json').read_text())
    author=json.loads((ROOT/'datasets/swe_bench_lite_gt_entities_3levels.json').read_text())
    projects=baseline_projects()
    common={name:common_metrics(p['records'],gold,ids) for name,p in projects.items()}
    common['GraphLocator/Qwen-适配v2']=dict(**metrics['common'],details=details)
    old_native=json.loads((ROOT/'results/qwen_300/full/loc_results.json').read_text())
    new_native=json.loads((folder/'loc_results.json').read_text())
    subset={}
    for name,iids in [('开发集',inputs['development']),('验证集',inputs['validation'])]:
        subset[name]={project:common_metrics(value['records'] if 'records' in value else records,gold,iids)
            for project,value in {**projects,'GraphLocator/Qwen-适配v2':{}}.items()}
    paired={name:paired_interval(common['GraphLocator/Qwen']['details'],details,iids)
        for name,iids in [('全部300题',ids),('验证集',inputs['validation'])]}
    pipeline,pipeline_details=analyze(folder,ids,gold,details,audits)
    development={}
    development_root=WORK/'development'/manifest['hash']
    for variant in ('B1','B2','B3'):
        scores,_,_=evaluate_folder(development_root/variant,inputs['development'])
        development[variant]=scores
    dataset=read_jsonl(ROOT/'datasets/swe_bench_lite.jsonl')
    by_project={}
    for repo in sorted({r['repo'] for r in dataset.values()}):
        iids=[iid for iid in ids if dataset[iid]['repo']==repo]
        by_project[repo]={name:common_metrics(p['records'] if 'records' in p else records,gold,iids)['strict']
            for name,p in {**projects,'GraphLocator/Qwen-适配v2':{}}.items()}
    deltas=[]
    for iid in ids:
        old=common['GraphLocator/Qwen']['details'][iid]['strict'];new=details[iid]['strict']
        if old is not None and new is not None:
            deltas.append(dict(instance_id=iid,old_f1=old['f1'],new_f1=new['f1'],delta=new['f1']-old['f1']))
    deltas.sort(key=lambda v:(-v['delta'],v['instance_id']))
    bundle=dict(common={name:{k:v for k,v in data.items() if k!='details'} for name,data in common.items()},
        subsets={name:{project:{k:v for k,v in data.items() if k!='details'} for project,data in values.items()} for name,values in subset.items()},
        paired=paired,development=development,pipeline=pipeline,by_project=by_project,per_issue_changes=deltas,
        baseline_project_sha256={name:project_hash(p) for name,p in projects.items()},version_hash=manifest['hash'],
        scoring_sha256={name:sha(ROOT/'qwen_adapt'/name) for name in ('evaluate.py','diagnostics.py','reports.py')})
    atomic(WORK/'comparison_metrics.json',bundle)
    atomic(WORK/'comparison_per_issue.json',{name:data['details'] for name,data in common.items()})
    overview='# GraphLocator + Qwen 适配v2：SWE-bench Lite 300题总览\n\n'
    overview+=f"300题均有状态记录：完成 {metrics['completed']}，失败 {metrics['failed']}。开发集30题与本轮未调参验证集270题分开报告。冻结版本 `{manifest['hash']}`；全部角色使用本地 Qwen2.5-Coder-7B-Instruct。\n\n"
    overview+=f"正式主条件为 **{selected}**。开发集B2共同严格函数F1 12.62%、原生函数F1 15.10%，高于B3的7.26%和8.20%，因此保留搜索与ID相关性适配，因果阶段保留原Mermaid协议。B3的中性JSON因果组合工程更稳，开发定位效果未支持采用，全部结果留档。B2仍存在旧因果上下文超限风险，实际失败计0。最终条件在270题验证前选定。\n\n"
    original=common['GraphLocator/Qwen']['strict'];adapted=metrics['common']['strict']
    overview+=f"共同严格函数F1由 {rate(original['f1'])} 变为 {rate(adapted['f1'])}，变化 {adapted['f1']-original['f1']:+.2%}；Precision变化 {adapted['precision']-original['precision']:+.2%}，Recall变化 {adapted['recall']-original['recall']:+.2%}。这描述本批固定结果，改善幅度以验证集配对区间判断；工程错误减少与原因判断变准分别看。\n\n"
    overview+='## 相同共同真值下的函数集合指标\n\n'
    overview+='| 项目/条件 | 适用题数 | Precision | Recall | 宏平均F1 | SL |\n|---|---:|---:|---:|---:|---:|\n'
    for name,data in common.items():
        overview+=metric_line(name,data['strict'])
    overview+='\n严格匹配先用 base commit 的 AST 识别完全限定函数名；裸函数名只在同文件唯一时映射。多义名称保留为未解析预测，不借助真值选类，仍占预测数量。金标的限定名依据原有 patch/AST 证据离线恢复。末段名集合指标另存 comparison_metrics.json，不与严格指标混用。\n\n'
    overview+='## 原生论文口径\n\n'
    overview+='| 层级/条件 | 适用题数 | Precision | Recall | 宏平均F1 | SL |\n|---|---:|---:|---:|---:|---:|\n'
    native_old=native_metrics(old_native,author,ids)
    for key,label in [('found_files','文件'),('found_modules','模块'),('found_functions','函数')]:
        overview+=metric_line(label+'/历史Qwen',native_old[key])
        overview+=metric_line(label+'/适配v2',metrics['native'][key])
    overview+='\n原生真值适用文件/模块/函数分别300/225/277题，保留作者已知字面格式问题。每题先算F1再平均；失败与空预测计0，无该层级真值计NA。原生集合不声称具有置信度排名。\n\n'
    overview+='## 开发与验证分开比较\n\n'
    overview+='| 子集/条件 | 适用题数 | Precision | Recall | 宏平均F1 | SL |\n|---|---:|---:|---:|---:|---:|\n'
    for split,values in subset.items():
        for name in ('GraphLocator/Qwen','GraphLocator/Qwen-适配v2'):
            overview+=metric_line(split+'/'+name,values[name]['strict'])
    for split,p in paired.items():
        if p:
            overview+=f"\n{split}严格函数F1配对变化：{p['mean']:+.2%}，逐题配对bootstrap 95%区间 [{p['low']:+.2%}, {p['high']:+.2%}]；改善/退化/相同为 {p['positive']}/{p['negative']}/{p['unchanged']}。区间来自这批题的重采样，不代表跨种子稳定性或全新benchmark泛化。\n"
    overview+='\n原生函数同题子集比较：\n\n| 子集/条件 | 适用题数 | Precision | Recall | 宏平均F1 | SL |\n|---|---:|---:|---:|---:|---:|\n'
    for split,iids in [('开发集',inputs['development']),('验证集',inputs['validation'])]:
        for label,predictions in [('历史Qwen',old_native),('适配v2',new_native)]:
            overview+=metric_line(split+'/'+label,native_metrics(predictions,author,iids)['found_functions'])
    overview+='\n## 历史表格兼容Top-k\n\n| 条件 | Top-1 | Top-3 | Top-5 |\n|---|---:|---:|---:|\n'
    for name,data in common.items():
        d=data['leaf']['denominator']
        overview+='| '+name+' | '+' | '.join(f"{data['topk'][k]}/{d}（{data['topk'][k]/d:.2%}）" for k in ('top1','top3','top5'))+' |\n'
    overview+='\n这里保留既有文件+函数末段名、类/变量占位及输出顺序的Top-k规则；不是作者原生排名指标。CoSIL/on/off复用其历史文件候选，GraphLocator重新检索全仓库，不能将分差归因于单个组件。\n\n'
    overview+='## 固定30题开发分项对照\n\n| 条件 | 完成/失败 | 严格函数F1 | 原生函数F1 | 空种子 | 源码送达缺口 | 实际模型请求 |\n|---|---:|---:|---:|---:|---:|---:|\n'
    for variant,m in development.items():
        overview+=f"| {variant} | {m['completed']}/{m['failed']} | {rate(m['common']['strict']['f1'])} | {rate(m['native']['found_functions']['f1'])} | {m['audit']['zero_seeds']} | {m['audit']['functions_without_source_receipt']} | {m['audit']['calls']} |\n"
    overview+='\nB1包含基础协议与搜索恢复；B2同时改变相关性输出格式和源码表示；B3再改变因果提示、协议及源码分批。B1/B2仍用旧CausalAgent的整批邻居输入，可能上下文超限；失败计0，未裁掉源码后伪装成功。因此B2→B3是因果适配组合，不能把效果全部归因于JSON格式。开发成绩不覆盖正式重新执行的30题，不据小样本断言单个改动的独立贡献。所有条件与历史B0的同题比较见开发审查记录。\n\n'
    overview+='## 按项目比较\n\n| 项目 | 函数适用题数 | 历史GraphLocator F1 | 适配v2 F1 | CoSIL/off F1 |\n|---|---:|---:|---:|---:|\n'
    for repo,values in by_project.items():
        overview+=f"| {repo} | {values['GraphLocator/Qwen']['denominator']} | {rate(values['GraphLocator/Qwen']['f1'])} | {rate(values['GraphLocator/Qwen-适配v2']['f1'])} | {rate(values['CoSIL/off']['f1'])} |\n"
    overview+='\n小项目样本少，单题可明显改变均值，不作项目层面的统计优越性宣称。\n\n'
    overview+='## 改善与退化案例\n\n| issue | 历史严格F1 | 适配v2严格F1 | 变化 |\n|---|---:|---:|---:|\n'
    shown=[d for d in deltas if d['delta']>0][:5]+sorted([d for d in deltas if d['delta']<0],key=lambda d:d['delta'])[:5]
    for d in shown:
        overview+=f"| {d['instance_id']} | {rate(d['old_f1'])} | {rate(d['new_f1'])} | {d['delta']:+.2%} |\n"
    overview+='\n案例按逐题F1差选择用于解释结果，未用于选择开发题或模型参数。具体目标、预测、阶段覆盖及模型原始因果依据见逐题分析；模型依据仅作原文留档，不自动视为正确归因。\n\n'
    overview+='## 阶段覆盖、选择损失与剩余问题\n\n'
    overview+='| 阶段 | 至少覆盖一个目标的题数 | 覆盖全部目标的题数 |\n|---|---:|---:|\n'
    for stage,label in [('search_received','搜索实际收到'),('seeds','搜索保留种子'),('causal_received','因果实际收到'),('final','最终预测')]:
        overview+=f"| {label} | {pipeline['counts'].get(stage+'_any_target',0)} | {pipeline['counts'].get(stage+'_all_targets',0)} |\n"
    overview+='\n'+pipeline['coverage_definition']+' 该统计是实体覆盖，不自动证明模型读懂源码。\n\n'
    for reason in ('运行失败','目标函数实体未送达模型','搜索收到目标实体但未进入因果','因果收到目标实体但未保留','最终覆盖部分目标','最终覆盖全部目标','NA'):
        overview+=f"- {reason}：{pipeline['counts'].get(reason,0)}题。\n"
    overview+=f"\n预测函数数：中位 {pipeline['predicted_function_count'].get('median',0)}、平均 {pipeline['predicted_function_count'].get('mean',0):.2f}、最多 {pipeline['predicted_function_count'].get('maximum',0)}；{pipeline['counts'].get('issues_with_duplicate_code_factors',0)}题有多个因子关联相同代码。重复原因和过多候选可能损害精确率，这一解释需结合实际P/R，不将其称为已证实的单一原因。\n\n"
    overview+=f"剩余协议问题：{pipeline['counts'].get('issues_with_unknown_relevance',0)}题存在相关性unknown；{pipeline['counts'].get('issues_with_unprocessed_matches',0)}题有预算内未处理搜索匹配；旧因果按解析调用计有 {pipeline['counts'].get('legacy_actual_update_limit',0)}题达到最多21次更新。旧Mermaid未执行B3的JSON校验，不能用causal_validation_errors=0声称因果输出全部合法。按原解析条件离线复查，有 {pipeline['counts'].get('issues_with_legacy_invalid_binding',0)}题、{pipeline['event_totals'].get('legacy_invalid_new_binding_rows',0)}条新增因子绑定行编号非法或越界；原程序丢弃这些行的代码绑定，保留原预测，未重新纠正后替换结果。\n\n"
    overview+=f"{pipeline['counts'].get('issues_with_causal_cycles',0)}题的最终图存在环；当前保持仓库有向图定义，未增加DAG约束。这一诊断不能单独证明定位失败，需结合各题源码与图解释。\n\n"
    overview+=f"代码关联中有 {pipeline['event_totals'].get('associations_without_file',0)}个因子指向无文件位置的图实体，{pipeline['event_totals'].get('associations_without_source',0)}个没有源码内容；{pipeline['counts'].get('issues_with_only_unlocated_associations',0)}题的关联全部没有文件位置。字段/变量和函数分别统计，图实体ID合法不等于获得可编辑函数定位；这可以造成种子与CIG非空而原生导出为空，逐题保留这些情况。\n\n"
    overview+='### 成本与概率\n\n'
    overview+=f"逐题耗时中位数 {pipeline['seconds'].get('median',0):.1f}秒、均值 {pipeline['seconds'].get('mean',0):.1f}秒；这些是并发运行的单题时间之和，不能当作墙钟耗时。实际调用和token如下。模型因果概率的分布见机器诊断；概率未经校准，不解释成定位成功率。\n\n"
    completion=WORK/'full/completion.json'
    wall=datetime.fromtimestamp(completion.stat().st_mtime).astimezone()-datetime.fromisoformat(manifest['started_at'])
    resources=manifest.get('resume_history',[])
    resource_text='初始四卡，用户要求释放GPU3后使用GPU0/1/2三卡，每卡两worker' if resources else '四卡每卡两worker'
    overview+=f"正式300题墙钟耗时约 {wall.total_seconds()/3600:.2f}小时；{resource_text}。因果边概率中位数 {pipeline['causal_probabilities'].get('median',0):.3f}、均值 {pipeline['causal_probabilities'].get('mean',0):.3f}。接近1的值只是模型判断，不当作有统计依据的因果概率。\n\n"
    resource_file=WORK/'resource_change.json'
    if resource_file.exists():
        resource=json.loads(resource_file.read_text())
        overview+=f"释放GPU3时保留 {resource['completed_cases_retained']}个已结束结果，{len(resource['interrupted_issues'])}个未完成尝试的事件、请求、日志移入interrupted_resource_change留档后重跑。已有预测字节哈希保留核验；未按定位分数挑选尝试。以上token与单题时间是最终采用尝试，资源切换的中断成本另存档，墙钟包含该过程。\n\n"
    overview+='## 行为与成本审计\n\n'
    overview+=json.dumps(metrics['audit'],ensure_ascii=False,indent=2).join(['```json\n','\n```\n\n'])
    overview+='实际送达按返回响应对应的请求及源码快照核验；文件名/类头不充当完整函数体。格式/编号合法不证明语义正确。删除具体示例后旧标签消失也不独立证明推理改进。\n\n'
    gaps=[(iid,a['final_functions_without_source_receipt']) for iid,a in audits.items() if a['final_functions_without_source_receipt']]
    if gaps:
        overview+='最终函数源码送达缺口：`'+json.dumps(gaps,ensure_ascii=False)+'`。该例缓存图函数源码为空，旧Mermaid仍可按编号选择此节点；空源码及其输出原样保留，没有把工具提取到名称当作收到函数代码，亦未在全量结束后补读源码重跑挑选结果。\n\n'
    overview+=f"运行失败率 {metrics['failed']/300:.2%}，相较旧40/300下降，但没有达到计划中争取低于5%的目标；完成题空种子率 {metrics['audit']['zero_seeds']/metrics['completed']:.2%}，也未达争取低于30%的目标。工程目标未达标与定位指标提升同时报告。\n\n"
    overview+=f"{pipeline['counts'].get('function_predictions_with_no_main_function_source',0)}题有函数预测而主搜索动作角色未收到显式函数体；{pipeline['counts'].get('function_predictions_with_no_causal_function_source',0)}题有函数预测而因果角色未收到显式函数体。相关性判别、主搜索动作与因果阶段的实际输入分开留档，不能用工具已提取源码推断所有角色都收到相同上下文；逐题分析列出三个角色的函数体数量。最终函数自身的源码来源仍单独核查。\n\n"
    overview+='## 复现中发现的问题与适配差异\n\n'
    overview+='正式B2采用经schema校验的普通JSON搜索动作、带ID的相关性判别及可读相关性源码；仍保留原Mermaid因果输出和整批邻居输入。搜索广泛查询最多4页，因果参数20按原实际最多21次更新对齐；原生因果输入超过65536时明确失败，未静默截断。稳定ID实现还修复了增量图同一类结束列相差1的错误碰撞。协议、查询顺序和搜索预算存在行为差异，属于Qwen适配实验。\n\n'
    overview+='开发B3额外采用中性因果提示、稳定ID JSON图、每个因子的causes边、关联替换与源码分批；效果未支持选为正式主条件。B3的语法合法性提升与函数F1退化同时报告，未因其代码改动更多而强行保留。\n\n'
    overview+='第一轮金丝雀虽完成却因未知因子操作得到空定位，未放行；后续轮次及代码快照完整保留。服务启动初期复用了原启动日志路径，随后转入独立目录，原逐题日志和预测未改动。详见 service_log_note.json 和各 canary/review.json。\n\n'
    overview+='另发现论文与仓库的优先级定义有差异：论文3.2节使用 `1-∏(1-p)`，原causal_agent.py实际按最大出边概率排序。本轮保留仓库的最大出边规则，未与模型适配同时更改，故不能称为论文公式的严格复现。稳定排序与分页也会改变探索顺序，效果不能只归因于JSON格式。\n\n'
    overview+=f'[机器指标](../../results/qwen_adapt_v2/full/{selected}/metrics.json) · [共同指标](../../results/qwen_adapt_v2/comparison_metrics.json) · [计划](GraphLocator_Qwen模型适配_问题与改善计划.md) · [开发与冻结审查](开发集对照与冻结审查_适配v2.md) · [逐题分析](实验报告_GraphLocator_Qwen适配v2_逐题分析.md)\n'
    detailed='# GraphLocator + Qwen 适配v2：逐题分析\n\n预测已冻结，以下为离线评分与日志审计。无函数真值为NA，失败和空预测分别显示。\n\n'
    for number,iid in enumerate(ids,1):
        row=records[iid];case_path=folder/('cases' if row['status']=='complete' else 'failures')/f'{iid}.json'
        case=json.loads(case_path.read_text());detail=details[iid];audit=audits[iid]
        detailed+=f"## {number}. {iid}\n\n项目：{dataset[iid]['repo']}；本轮子集：{'开发集' if iid in inputs['development'] else '验证集'}；状态：{row['status']}。\n\n"
        detailed+='问题描述节选：\n\n'+'\n'.join('> '+line for line in dataset[iid]['problem_statement'][:800].splitlines())+'\n\n'
        detailed+='| 条件 | 严格P | 严格R | 严格F1 | SL | Top-5 |\n|---|---:|---:|---:|---:|---:|\n'
        for name,data in common.items():
            result=data['details'][iid];m=result['strict']
            detailed+='| '+name+' | '+' | '.join(rate(m[k]) if m else 'NA' for k in ('precision','recall','f1','sl'))+' | '+('NA' if result['top5'] is None else '命中' if result['top5'] else '未命中')+' |\n'
        detailed+='\n共同严格目标：`'+json.dumps(detail['strict_targets'],ensure_ascii=False)+'`\n\n'
        detailed+='新预测：`'+json.dumps(detail['strict_predictions'],ensure_ascii=False)+'`\n\n'
        detailed+=f"搜索种子 {len(case.get('seed_locations',[]))}；搜索停止={case.get('search_stop_reason','未完成')}；因果停止={case.get('causal_stop_reason','未完成')}；输入/输出token={audit['tokens'].get('input',0)}/{audit['tokens'].get('output',0)}；耗时={case.get('seconds',0):.1f}秒。\n\n"
        stage=pipeline_details[iid]
        detailed+='阶段诊断：'+stage['reason']+'；各阶段实际覆盖目标：`'+json.dumps(stage['coverage'],ensure_ascii=False)+'`。\n\n'
        detailed+='各角色实际收到显式函数体数：`'+json.dumps(stage['source_function_counts_by_role'],ensure_ascii=False)+'`。\n\n'
        if row['status']=='failed':
            detailed+='故障：'+str(case.get('error'))+'；'+str(case.get('message'))[:500]+'\n\n'
        if detail['unresolved']:
            detailed+='无法唯一映射的函数：`'+json.dumps(detail['unresolved'],ensure_ascii=False)+'`\n\n'
        for factor in case.get('causal_graph',{}).get('nodes',[]):
            if factor['id']=='I':continue
            detailed+='- '+factor['id']+'：'+str(factor.get('name',''))+'；依据：'+str(factor.get('evidence',''))+'；关联：`'+json.dumps([(n.get('file'),n.get('qualified_name')) for n in factor.get('code_elements',[])],ensure_ascii=False)+'`\n'
        detailed+='\n源码送达缺口：`'+json.dumps(audit['final_functions_without_source_receipt'],ensure_ascii=False)+'`；相关性未知项='+str(audit['unknown_relevance'])+'；因果校验错误项='+str(audit['causal_validation_errors'])+'。\n\n'
        detailed+='[原始结果](../../results/qwen_adapt_v2/full/'+selected+'/'+('cases' if row['status']=='complete' else 'failures')+'/'+iid+'.json) · [事件日志](../../results/qwen_adapt_v2/full/'+selected+'/events/'+iid+'.jsonl)\n\n'
    (ROOT/'docs/qwen-300/实验报告_GraphLocator_Qwen适配v2_总览.md').write_text(overview)
    (ROOT/'docs/qwen-300/实验报告_GraphLocator_Qwen适配v2_逐题分析.md').write_text(detailed)
    return bundle,common


def supplement_workbook(bundle,common):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.path.insert(0,str(COSIL/'code_v3b/.matrix_deps'))
        from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill
    path=COSIL/'docs/qwen-300/跨实验逐题定位对照表.xlsx'
    book=load_workbook(path)
    for name in ['集合指标_适配v2','逐题集合F1_适配v2']:
        if name in book:
            del book[name]
    sheet=book.create_sheet('集合指标_适配v2')
    sheet.append(['评分口径','项目/条件','适用题数','Precision','Recall','宏平均F1','SL'])
    for name,data in common.items():
        for scope in ('strict','leaf'):
            m=data[scope];sheet.append([scope,name,m['denominator'],*[m[k] for k in ('precision','recall','f1','sl')]])
    per=book.create_sheet('逐题集合F1_适配v2')
    names=list(common)
    per.append(['issue ID','子集',*names])
    inputs=json.loads((WORK/'inputs_manifest.json').read_text())
    for iid in inputs['instances']:
        per.append([iid,'开发集' if iid in inputs['development'] else '验证集',
            *[common[n]['details'][iid]['strict']['f1'] if common[n]['details'][iid]['strict'] else 'NA' for n in names]])
    for ws in (sheet,per):
        ws.freeze_panes='C2'
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='305496')
        for row in ws.iter_rows(min_row=2):
            for cell in row[3 if ws==sheet else 2:]:
                if isinstance(cell.value,(int,float)):
                    cell.number_format='0.00%'
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width=30
    book.save(path)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--workbook',action='store_true')
    args=parser.parse_args()
    bundle,common=generate()
    if args.workbook:
        supplement_workbook(bundle,common)
    print(json.dumps(bundle,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
